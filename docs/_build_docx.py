from docx import Document
from docx.shared import Pt, Inches, RGBColor
from pathlib import Path
import re


def apply_base_style(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for lvl, size in [(1, 18), (2, 14), (3, 12)]:
        heading = doc.styles[f"Heading {lvl}"]
        heading.font.name = "Calibri"
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(0x14, 0x21, 0x3D)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)


CODE_FONT = "Consolas"


def add_inline(paragraph, text: str):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(10)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def flush_paragraph(doc, buffer):
    if not buffer:
        return
    text = " ".join(line.strip() for line in buffer).strip()
    if text:
        p = doc.add_paragraph()
        add_inline(p, text)
    buffer.clear()


def md_to_docx(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    apply_base_style(doc)

    in_code = False
    code_buffer: list[str] = []
    para_buffer: list[str] = []
    list_buffer: list[str] = []

    def flush_list():
        for item in list_buffer:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, item)
        list_buffer.clear()

    for line in lines:
        if line.startswith("```"):
            flush_paragraph(doc, para_buffer)
            flush_list()
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buffer))
                run.font.name = CODE_FONT
                run.font.size = Pt(10)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_buffer.append(line)
            continue

        stripped = line.strip()

        if stripped.startswith("### "):
            flush_paragraph(doc, para_buffer)
            flush_list()
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            flush_paragraph(doc, para_buffer)
            flush_list()
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            flush_paragraph(doc, para_buffer)
            flush_list()
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            flush_paragraph(doc, para_buffer)
            list_buffer.append(stripped[2:])
        elif stripped == "":
            flush_paragraph(doc, para_buffer)
            flush_list()
        else:
            if list_buffer and line.startswith("  "):
                list_buffer[-1] += " " + stripped
            else:
                flush_list()
                para_buffer.append(stripped)

    flush_paragraph(doc, para_buffer)
    flush_list()
    doc.save(out_path)


def main():
    docs_dir = Path(__file__).parent
    conversions = [
        ("production-thinking.md", "Part 2 - Production Thinking.docx"),
        ("frontend-architecture.md", "Part 3 - Frontend Architecture.docx"),
    ]
    for src, dst in conversions:
        md_to_docx(docs_dir / src, docs_dir / dst)
        print(f"Wrote {dst}")


if __name__ == "__main__":
    main()
